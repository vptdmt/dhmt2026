import sys
import openpyxl
import json

sys.stdout.reconfigure(encoding='utf-8')

# Load existing data.json to preserve categories and other tabs
with open('data.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

old_runners_map = {r['bib']: r for r in data['funrun']['runners']}

wb = openpyxl.load_workbook('Danh sach xep xe Fun Run_19.10.xlsx', data_only=True)
s1 = wb['Tổng hợp điều xe Fun Run']
s2 = wb['Chi tiết danh sách từng xe']

# 1. Parse Sheet 1 (Tổng hợp 33 xe)
sheet1_vehicles = []
for r in list(s1.iter_rows(values_only=True))[4:]:
    if r[0] is not None and str(r[0]).strip().isdigit():
        num = int(r[0])
        name = str(r[1]).strip() if r[1] else f'Xe {num:02d}'
        cap_str = str(r[2]).strip() if r[2] else '7 chỗ'
        pax_str = str(r[3]).strip() if r[3] else ''
        depart = str(r[4]).strip() if r[4] else '05:00'
        route = str(r[5]).strip() if r[5] else 'Sảnh Khách sạn Mường Thanh Luxury Ninh Bình ➔ Sân Golf Tràng An & Resort'
        captain = str(r[6]).strip() if r[6] else ''
        captain_phone = str(r[7]).strip() if r[7] else ''
        driver = str(r[8]).strip() if r[8] else ''
        driver_phone = str(r[9]).strip() if r[9] else ''
        bib_range = str(r[10]).strip() if r[10] else ''
        plate = str(r[11]).strip() if r[11] else ''
        note = str(r[12]).strip() if r[12] else ''

        # Capacity number
        cap_num = 7
        for c in [45, 29, 16, 7, 4]:
            if str(c) in cap_str:
                cap_num = c
                break
        if 'Tự túc' in cap_str or num in [30, 31]:
            if num == 30: cap_num = 13
            elif num == 31: cap_num = 15

        sheet1_vehicles.append({
            'num': num,
            'id': f'XE-{num:02d}',
            'name': name,
            'type': cap_str,
            'capacity': cap_num,
            'paxSummary': pax_str,
            'departTime': depart,
            'route': route,
            'captain': captain,
            'captainPhone': captain_phone,
            'driver': driver,
            'driverPhone': driver_phone,
            'bibRange': bib_range,
            'plate': plate,
            'note': note
        })

print(f'Parsed {len(sheet1_vehicles)} vehicles from Sheet 1.')

# 2. Parse Sheet 2 (Chi tiết danh sách từng xe)
sheet2_data = {}
cur_xe_num = None
for r in s2.iter_rows(values_only=True):
    val0 = str(r[0]).strip() if r[0] is not None else ''
    if val0.startswith('XE '):
        parts = val0.split('/')
        xe_part = parts[0].replace('XE', '').strip()
        try:
            cur_xe_num = int(xe_part)
        except:
            cur_xe_num = val0
        sheet2_data[cur_xe_num] = {'header': val0, 'members': []}
        continue
    if cur_xe_num and val0 in ['STT', '']:
        continue
    if cur_xe_num and r[2] is not None:
        raw_bib = r[1]
        bib_val = None
        is_btc = False
        if raw_bib is not None:
            raw_bib_str = str(raw_bib).strip()
            if raw_bib_str.isdigit():
                bib_val = int(raw_bib_str)
            elif raw_bib_str.upper() == 'BTC':
                is_btc = True

        name = str(r[2]).strip()
        gender = str(r[3]).strip() if r[3] else ''
        title = str(r[4]).strip() if r[4] else ''
        unit = str(r[5]).strip() if r[5] else ''
        dist = str(r[6]).strip() if r[6] else ''
        size = str(r[7]).strip() if r[7] else ''
        phone = str(r[8]).strip() if r[8] else ''
        role = str(r[9]).strip() if r[9] else ''

        sheet2_data[cur_xe_num]['members'].append({
            'stt': r[0],
            'bib': bib_val,
            'isBTC': is_btc,
            'name': name,
            'gender': gender,
            'title': title,
            'unit': unit,
            'dist': dist,
            'size': size,
            'phone': phone,
            'role': role
        })

print(f'Parsed members for {len(sheet2_data)} vehicles from Sheet 2.')

# 3. Assemble funrun.vehicles
vehicles = []
all_runners = []
total_runners_count = 0

for v1 in sheet1_vehicles:
    num = v1['num']
    v_id = v1['id']
    s2_info = sheet2_data.get(num, {'header': '', 'members': []})
    members = s2_info['members']
    
    # Identify captain title & unit
    cap_name = v1['captain']
    cap_phone = v1['captainPhone']
    cap_title = ''
    cap_unit = ''
    
    # Try finding captain in members
    for m in members:
        if m['name'] and cap_name and (m['name'].lower() in cap_name.lower() or cap_name.lower() in m['name'].lower()):
            cap_title = m['title']
            cap_unit = m['unit']
            if not cap_phone and m['phone']:
                cap_phone = m['phone']
            break

    # If captain still missing title/unit, search in old btcStaff
    if not cap_title or not cap_unit:
        for b in data['funrun'].get('btcStaff', []):
            if b['name'] and cap_name and b['name'].lower() in cap_name.lower():
                cap_title = b.get('title', '')
                cap_unit = b.get('dept', '')
                if not cap_phone and b.get('phone'):
                    cap_phone = b['phone']
                break

    # Vehicle label
    if num in [1, 2, 3, 4, 30, 31, 32, 33]:
        v_label = f"Xe {num:02d} - {v1['name']}"
    else:
        v_label = f"Xe {num:02d} - {v1['name']} ({v1['type']})"

    v_obj = {
        'id': v_id,
        'num': num,
        'code': f'XE-{num:02d}',
        'label': v_label,
        'type': v1['type'],
        'capacity': v1['capacity'],
        'paxCount': len(members),
        'paxSummary': v1['paxSummary'],
        'departTime': v1['departTime'],
        'pickupLoc': v1['route'].split('➔')[0].strip() if '➔' in v1['route'] else 'Sảnh Khách sạn Mường Thanh Luxury Ninh Bình',
        'destLoc': 'Sân Golf Tràng An & Resort',
        'returnTime': '11:00',
        'captain': {
            'name': cap_name,
            'title': cap_title,
            'unit': cap_unit,
            'phone': cap_phone
        },
        'driver': {
            'name': v1['driver'],
            'phone': v1['driverPhone']
        },
        'plate': v1['plate'],
        'bibRange': v1['bibRange'],
        'note': v1['note'],
        'isVIP': (num in [1, 2]),
        'isEarly': (num == 3 or '04:30' in str(v1['departTime'])),
        'isEventTeam': (num == 4),
        'isSelfTransport': (num in [30, 31]),
        'isMedical': (num == 32),
        'isCoordination': (num == 33),
        'members': [],
        'runners': [],
        'btcStaff': [],
        'drivers': []
    }

    for m in members:
        mem_obj = {
            'stt': m['stt'],
            'bib': m['bib'],
            'name': m['name'],
            'gender': m['gender'],
            'title': m['title'],
            'unit': m['unit'],
            'dist': m['dist'],
            'size': m['size'],
            'phone': m['phone'],
            'role': m['role'],
            'vehicleId': v_id,
            'vehicleNum': num,
            'vehicleLabel': v_label,
            'captainName': cap_name,
            'captainPhone': cap_phone,
            'driverName': v1['driver'],
            'driverPhone': v1['driverPhone'],
            'plate': v1['plate']
        }
        v_obj['members'].append(mem_obj)

        if m['bib'] is not None:
            # Runner
            # Determine category
            bib = m['bib']
            old_r = old_runners_map.get(bib, {})
            cat = old_r.get('category', '')
            if not cat:
                # Infer category from dist & gender
                d_str = m['dist'].lower()
                g_str = m['gender']
                if '15' in d_str: cat = f"15km {g_str}"
                elif '10' in d_str: cat = f"10km {g_str}"
                elif '5' in d_str: cat = f"5km {g_str}"
                else: cat = '5km'

            runner_dist = '5km'
            if '10' in m['dist']: runner_dist = '10km'
            elif '15' in m['dist']: runner_dist = '15km'

            is_btc_role = ('BTC' in m['role'] or 'Trưởng xe' in m['role'] or old_r.get('btc'))
            btc_badge = m['role'] if ('BTC' in m['role'] or 'Trưởng xe' in m['role']) else (old_r.get('btc', '') if is_btc_role else '')

            runner_obj = {
                'stt': len(all_runners) + 1,
                'bib': bib,
                'name': m['name'],
                'gender': m['gender'],
                'title': m['title'],
                'unit': m['unit'],
                'distance': runner_dist,
                'dist': m['dist'],
                'category': cat,
                'size': m['size'] or old_r.get('size', 'M'),
                'phone': m['phone'] or old_r.get('phone', ''),
                'role': m['role'],
                'btc': btc_badge,
                'note': '',
                'vehicleId': v_id,
                'vehicleNum': num,
                'vehicleLabel': v_label,
                'captainName': cap_name,
                'captainPhone': cap_phone,
                'captainTitle': cap_title,
                'captainUnit': cap_unit,
                'driverName': v1['driver'],
                'driverPhone': v1['driverPhone'],
                'plate': v1['plate']
            }
            v_obj['runners'].append(runner_obj)
            all_runners.append(runner_obj)
        elif m['isBTC']:
            v_obj['btcStaff'].append({
                'name': m['name'],
                'gender': m['gender'],
                'title': m['title'],
                'unit': m['unit'],
                'phone': m['phone'],
                'role': m['role']
            })
        else:
            v_obj['drivers'].append({
                'name': m['name'],
                'gender': m['gender'],
                'title': m['title'],
                'phone': m['phone'],
                'role': m['role']
            })

    vehicles.append(v_obj)

all_runners.sort(key=lambda r: r['bib'])
for idx, r in enumerate(all_runners):
    r['stt'] = idx + 1

print(f'Total vehicles built: {len(vehicles)}')
print(f'Total runners with BIB built: {len(all_runners)}')

# Update funrun in data
data['funrun']['vehicles'] = vehicles
data['funrun']['runners'] = all_runners
data['funrun']['overview']['totalRunners'] = len(all_runners)
data['funrun']['overview']['totalVehicles'] = len(vehicles)

# Recompute summary counts
c5 = [r for r in all_runners if r['distance'] == '5km']
c10 = [r for r in all_runners if r['distance'] == '10km']
c15 = [r for r in all_runners if r['distance'] == '15km']

data['funrun']['overview']['summaryCounts'] = {
    'total': len(all_runners),
    'c5k': len(c5),
    'c5k_m_lt35': sum(1 for r in c5 if r['category'] == '5km Nam dưới 35 tuổi'),
    'c5k_m_gte35': sum(1 for r in c5 if r['category'] == '5km Nam trên 35 tuổi'),
    'c5k_f_lt35': sum(1 for r in c5 if r['category'] == '5km Nữ dưới 35 tuổi'),
    'c5k_f_gte35': sum(1 for r in c5 if r['category'] == '5km Nữ trên 35 tuổi'),
    'c10k': len(c10),
    'c10k_m': sum(1 for r in c10 if r['gender'] == 'Nam'),
    'c10k_f': sum(1 for r in c10 if r['gender'] == 'Nữ'),
    'c15k': len(c15),
    'c15k_m': sum(1 for r in c15 if r['gender'] == 'Nam'),
    'c15k_f': sum(1 for r in c15 if r['gender'] == 'Nữ')
}

with open('data.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print('Updated data.json successfully!')
